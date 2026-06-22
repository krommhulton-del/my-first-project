# -*- coding: utf-8 -*-
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Inches as PInches, Pt as PPt, Emu
from pptx.dml.color import RGBColor as PColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

BASE = os.path.dirname(os.path.abspath(__file__))
OUT  = os.path.join(BASE, "策划大作业_纸上花火")
IMG  = os.path.join(OUT, "效果图")
imgs = {
    "room": os.path.join(IMG, "01_关键视觉_卧室黄昏.png"),
    "qq":   os.path.join(IMG, "02_WorldSpaceUI_QQ空间.png"),
    "mag":  os.path.join(IMG, "03_花火杂志_竖排金句.png"),
    "end":  os.path.join(IMG, "04_结局_褪色光化.png"),
}
FONT = "微软雅黑"
SERIF = "思源宋体"

# =========================================================== WORD
def cjk(run, name=FONT, size=10.5, bold=False, color=None):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {}); rPr.insert(0, rFonts)
    rFonts.set(qn('w:eastAsia'), name)

def para(doc, text, size=10.5, bold=False, color=None, align=None, after=4, name=FONT):
    p = doc.add_paragraph();
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text); cjk(r, name, size, bold, color)
    return p

def heading(doc, text, size=15, color=RGBColor(0x7A,0x2A,0x35)):
    para(doc, text, size=size, bold=True, color=color, after=6, name=SERIF)

def bullets(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style=None)
        p.paragraph_format.left_indent = Pt(14); p.paragraph_format.space_after = Pt(2)
        r = p.add_run("· " + it); cjk(r, FONT, size)

def table(doc, rows):
    t = doc.add_table(rows=len(rows), cols=len(rows[0])); t.style = 'Table Grid'
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            c = t.cell(i, j); c.paragraphs[0].clear()
            r = c.paragraphs[0].add_run(cell); cjk(r, FONT, 9.5, bold=(i == 0))
    doc.add_paragraph()

def build_word():
    doc = Document()
    # title
    para(doc, "《纸上花火 · Paper Fireworks》", size=24, bold=True,
         color=RGBColor(0x7A,0x2A,0x35), align=WD_ALIGN_PARAGRAPH.CENTER, after=2, name=SERIF)
    para(doc, "古早怀旧主题 · 沉浸式 VR 体验 ｜ 策划大作业（含 甲方需求文档 + 乙方解决方案）",
         size=11, color=RGBColor(0x80,0x70,0x60), align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, "（自己一组）", size=10, color=RGBColor(0xA0,0x90,0x80),
         align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

    # ---------- PART 1 需求文档 ----------
    para(doc, "第一部分　需求文档（甲方 / Party A）", size=16, bold=True,
         color=RGBColor(0x2C,0x3A,0x47), after=8, name=SERIF)
    para(doc, "说明：甲方可对功能模块自由定制，不考虑实现成本、经费与时间。本部分即「我想要一个什么样的产品」。",
         size=9.5, color=RGBColor(0x90,0x80,0x70), after=8)

    heading(doc, "一、项目概述")
    table(doc, [
        ["项", "内容"],
        ["产品名称", "《纸上花火 · Paper Fireworks》"],
        ["产品类型", "古早怀旧主题 · 沉浸式 VR 体验 / 叙事探索轻游戏"],
        ["目标平台", "PC VR（Meta Quest / Pico 串流），另出桌面体验版"],
        ["体验时长", "单次 15–25 分钟，可重复"],
        ["一句话概念", "戴上头显，回到 2005–2012 那个 QQ 空间、非主流、花火杂志的青春年代——在一间被时光封存的少女卧室里，翻动每一件旧物，唤醒一段集体记忆，完成一次对青春的温柔告别。"],
    ])
    heading(doc, "二、立意与情感内核")
    bullets(doc, [
        "母题：美好转瞬即逝；每个人的青春，都值得被温柔地收藏一次。",
        "它不是「通关取胜」的游戏，而是一座可以走进去的时代记忆博物馆。",
        "体验关键词：怀旧 · 治愈 · 沉浸 · 温柔的告别。",
    ])
    heading(doc, "三、目标用户与使用场景")
    bullets(doc, [
        "核心用户：90 后 / 95 后——QQ 空间原住民、花火/飞言情读者、非主流时代亲历者。",
        "扩展用户：Y2K 千禧美学爱好者、情感治愈与沉浸体验玩家、文创与展览观众。",
        "使用场景：个人头显体验 / 线下怀旧主题快闪展览 / 品牌联名（汽水·文具·音乐 App）。",
    ])
    heading(doc, "四、需求功能模块（甲方提出，不限可行性）")
    para(doc, "4.1 场景：一间被时光封存的千禧少女卧室", bold=True, size=11)
    para(doc, "高度还原 2005–2012 细节：碎花床单、贴满海报与明星贴纸的墙、老式台式电脑（CRT+QQ）、卡带机/MP3、梳妆台、铁皮文具盒、毛绒玩具、星空夜光顶贴、铁丝相片墙、成摞花火杂志；窗外是夏日傍晚的晚霞与蝉鸣、慢慢转的电风扇。")
    para(doc, "4.2 互动旧物（每件 = 一段可唤醒的记忆）", bold=True, size=11)
    table(doc, [
        ["旧物", "交互", "唤起的记忆"],
        ["老式台式机", "开机→进入 QQ 空间（闪图、非主流签名、自动播放空间音乐、踩空间）", "那年的网络青春"],
        ["卡带机/MP3", "放入磁带/戴耳机→播放那年的歌（原创致敬版避版权）", "一首歌一段心事"],
        ["花火/飞言情杂志", "抓取→翻页，古早言情插画 + 竖排金句", "课桌下偷看的言情"],
        ["同学录", "翻开→手写体留言「给十年后的你」", "毕业季"],
        ["钢笔+信纸", "写一封寄不出的信（手柄书写/语音），封存或点燃", "没说出口的话"],
        ["拍立得相机", "取景→按快门→生成做旧老照片贴上相片墙", "把此刻收藏"],
    ])
    para(doc, "4.3 核心玩法", bold=True, size=11)
    bullets(doc, ["第一人称 VR 漫游 + 手柄交互；", "唤醒旧物→收集「记忆碎片」→集齐解锁结局；",
                  "写信告别 + 拍照留念 两个情感仪式；", "无失败、无倒计时、治愈向，节奏舒缓。"])
    para(doc, "4.4 UI（两种模式都要）：Overlay（系统/进度/相册）+ WorldSpace（QQ 空间、杂志页贴在物体上）。", size=10.5)
    para(doc, "4.5 音乐音效：卡带底噪、蝉鸣、电扇、键盘、QQ 提示音、拨号上网音 + 原创年代主题曲。", size=10.5)
    para(doc, "4.6 粒子氛围：夕照浮尘、飘动窗帘、夜晚萤火/远处烟花、结局褪色光化。", size=10.5)
    para(doc, "4.7 结局：集齐记忆后房间在夕阳里褪色、光化、飘散，浮出落款金句，呼应母题。", size=10.5)
    heading(doc, "五、核心体验流程")
    para(doc, "进入房间（黄昏）→ 自由探索·逐件唤醒旧物 → 集齐记忆碎片 → 写一封寄不出的信 → 拍一张留念照 → 房间褪色光化·告别 → 落款金句。", size=10.5)
    heading(doc, "六、验收 / 效果期望")
    bullets(doc, ["还原度：目标用户一进门脱口而出「这就是我那时候的房间」。", "情感共鸣：至少一件旧物让人鼻子一酸。",
                  "沉浸与舒适：无眩晕、可久待。", "可传播：每帧可截图，适合展览/联名。"])

    doc.add_page_break()

    # ---------- PART 2 解决方案 ----------
    para(doc, "第二部分　解决方案（乙方 / Party B）", size=16, bold=True,
         color=RGBColor(0x2C,0x3A,0x47), after=8, name=SERIF)
    para(doc, "说明：本部分对甲方需求给出解决办法与「效果呈现的参考图片」，即「我要怎么做出来、做出来长什么样」。",
         size=9.5, color=RGBColor(0x90,0x80,0x70), after=8)

    heading(doc, "一、对需求的理解")
    para(doc, "甲方要的不是「游戏」，是一台情绪时光机——一切设计都围绕：让人一进去就想起自己的青春，并舍不得离开。还原度、声音、光与字迹是命门。")
    heading(doc, "二、设计总览（怎么解决）")
    table(doc, [
        ["甲方需求", "解决办法"],
        ["高还原千禧卧室", "3D 建模 + PBR 做旧材质 + 电影级怀旧后处理（暖调+暗角+颗粒+泛黄），把「塑料 3D」压成「旧照片」"],
        ["旧物 = 记忆", "每件旧物 = 可交互体（XR 抓取/触发）+ 触发演出（音乐/动画/竖排文字）+ 一片记忆碎片"],
        ["自由漫游不眩晕", "XR Interaction Toolkit；瞬移 + 舒适转向 + 暗角减晕；无失败无倒计时"],
        ["两种 UI", "Overlay 放系统/进度/相册；WorldSpace 把 QQ 空间、杂志页贴在物体上"],
        ["怀旧音乐避版权", "原创致敬版 BGM + 真实环境音，彻底规避版权"],
        ["治愈的告别", "结局粒子演出：房间褪色→光化→飘散 + 落款金句"],
    ])
    heading(doc, "三、美术与场景方案")
    bullets(doc, [
        "风格定位：千禧 Y2K × 古早言情 × 非主流——暖黄夕照、玫红网格、闪光渐变、衬线竖排。",
        "场景：约 12㎡ 少女卧室单场景（满足「一个关卡」），靠道具密度而非场景大小堆沉浸；窗外天幕做黄昏→入夜光照变化。",
        "色板：暖黄 #E8C36A / 玫红 #D96E8C / 米白 #E8DFD3 / 夜蓝 #2C3A47 / 高光白 #F5EBD8，统一压旧。",
    ])
    heading(doc, "四、互动设计方案（XR 手柄）")
    table(doc, [
        ["旧物", "交互方式", "触发演出"],
        ["台式机", "按电源→射线点 QQ 图标", "WorldSpace QQ 空间：闪图头像、非主流签名、自动播放、踩一踩"],
        ["卡带机", "抓磁带放入/戴耳机", "原创年代曲淡入 + 竖排歌词浮现"],
        ["杂志", "抓取 + 翻页（物理）", "古早言情插画 + 竖排金句逐行浮入"],
        ["信纸钢笔", "手柄书写/语音→封口/点燃", "信化作光点/灰烬飘出窗外 + 柔和音效"],
        ["拍立得", "举起取景 + 按快门", "咔嚓 + 吐出做旧照片贴上相片墙"],
    ])
    heading(doc, "五、系统与玩法方案")
    bullets(doc, ["记忆收集系统：8 件核心旧物 = 8 片记忆，集齐触发结局。",
                  "情感仪式系统：写信 + 拍照，结果存入相册/漂流瓶可回看。",
                  "软引导：未唤醒旧物带微光 + 轻响；无硬性引导。", "无失败设计：纯探索治愈，照顾展览观众。"])
    heading(doc, "六、叙事流程（节拍）")
    para(doc, "黄昏入场 → 软引导探索·逐件唤醒 → 集齐 8 片（推到入夜）→ 写信 → 拍照 → 房间褪色光化·旧物飘散 → 黑场落款：「那年夏天，我们都以为来日方长。」", size=10.5)
    heading(doc, "七、UI 方案（两模式落地）")
    bullets(doc, ["Overlay：进入标题卡 / 记忆进度（8 颗星）/ 相册回看 / 设置。",
                  "WorldSpace：CRT 里的 QQ 空间 / 杂志同学录页面 / 旧物悬浮提示。"])
    heading(doc, "八、音频方案")
    para(doc, "原创主题曲（钢琴+卡带底噪）+ 环境层（蝉鸣/电扇/拨号/QQ 音效）+ 交互音效。全部原创或免版税，彻底规避版权。", size=10.5)
    heading(doc, "九、技术方案")
    bullets(doc, ["引擎：Unity + URP（Bloom/暖调/颗粒/暗角）。",
                  "交互：XR Interaction Toolkit（Ray + Direct）+ XR Device Simulator（无头显可演示）。",
                  "文字：TextMeshPro 竖排 + 中文衬线。", "输出：PC VR 包 + 桌面演示版 + 360° 截图。"])
    heading(doc, "十、防眩晕与体验")
    para(doc, "瞬移移动 + 舒适转向 + 移动时暗角；交互在手边、视线舒适区；无快速位移与强闪烁；可坐可站。", size=10.5)

    heading(doc, "十一、效果呈现参考图（概念效果图）")
    cap = ["▲ 关键视觉 / 卧室黄昏", "▲ WorldSpace UI / 开机进 QQ 空间",
           "▲ 互动旧物 / 花火杂志翻页·竖排", "▲ 结局 / 褪色光化·告别"]
    for key, c in zip(["room", "qq", "mag", "end"], cap):
        if os.path.exists(imgs[key]):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(imgs[key], width=Inches(5.6))
            para(doc, c, size=9, color=RGBColor(0x80,0x70,0x60), align=WD_ALIGN_PARAGRAPH.CENTER, after=10)

    heading(doc, "十二、开发里程碑")
    table(doc, [["阶段", "内容"],
                ["P0 原型", "单房间 + 漫游 + 1 件旧物交互跑通"],
                ["P1 美术", "全场景道具 + 做旧材质 + 后处理调色"],
                ["P2 交互", "8 件旧物 + 记忆系统 + 写信/拍照"],
                ["P3 音频叙事", "BGM + 环境音 + 叙事流 + 结局演出"],
                ["P4 打磨", "防晕优化 + 桌面版 + 录屏与展览物料"]])
    heading(doc, "十三、风险与对策")
    table(doc, [["风险", "对策"],
                ["年代音乐版权", "全用原创致敬版 / 免版税"],
                ["VR 眩晕", "瞬移 + 舒适转向 + 暗角"],
                ["还原度不够「对味」", "海量真实道具参考 + 老用户测试"],
                ["体验偏静态", "加入写信/拍照两个主动仪式"]])
    heading(doc, "十四、商业与展览价值")
    para(doc, "怀旧是强情绪流量；天然适合线下快闪/展览与品牌联名（汽水、文具、音乐 App），每帧可截图传播，长尾与话题性强。", size=10.5)

    path = os.path.join(OUT, "纸上花火_策划大作业_需求+解决方案.docx")
    doc.save(path); print("WORD:", path)

# =========================================================== PPT
WINE = PColor(0x7A,0x2A,0x35); IVORY = PColor(0xE8,0xDF,0xD3)
GOLD = PColor(0xB8,0x96,0x64); SEAL = PColor(0xF5,0xEB,0xD8); BG = PColor(0x16,0x10,0x1A)

def pfont(run, name=FONT, size=18, color=IVORY, bold=False):
    f = run.font; f.size = PPt(size); f.bold = bold; f.color.rgb = color; f.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ('a:latin', 'a:ea', 'a:cs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set('typeface', name)

def build_ppt():
    prs = Presentation(); prs.slide_width = PInches(13.333); prs.slide_height = PInches(7.5)
    blank = prs.slide_layouts[6]
    W, H = prs.slide_width, prs.slide_height

    def slide(title, bullets_, img=None, notes="", big=False, sub=None):
        s = prs.slides.add_slide(blank)
        s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
        # accent line
        ln = s.shapes.add_textbox(PInches(0.6), PInches(1.35), PInches(3), PInches(0.05))
        # title
        tb = s.shapes.add_textbox(PInches(0.7), PInches(0.5), W - PInches(1.4), PInches(1.0))
        tf = tb.text_frame; tf.word_wrap = True
        r = tf.paragraphs[0].add_run(); r.text = title; pfont(r, SERIF, 34 if big else 28, GOLD, True)
        if sub:
            ps = tf.add_paragraph(); rs = ps.add_run(); rs.text = sub; pfont(rs, FONT, 16, IVORY)
        # image
        if img and os.path.exists(img):
            pic_w = PInches(6.6)
            s.shapes.add_picture(img, W - pic_w - PInches(0.6), PInches(1.7), width=pic_w)
            box_w = W - pic_w - PInches(1.8)
        else:
            box_w = W - PInches(1.6)
        # bullets
        if bullets_:
            bb = s.shapes.add_textbox(PInches(0.8), PInches(1.8), box_w, PInches(5.0))
            btf = bb.text_frame; btf.word_wrap = True
            for i, b in enumerate(bullets_):
                p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
                p.space_after = PPt(10)
                r = p.add_run(); r.text = "· " + b; pfont(r, FONT, 17, IVORY)
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    # P1 cover
    s = prs.slides.add_slide(blank); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    if os.path.exists(imgs["room"]):
        s.shapes.add_picture(imgs["room"], 0, 0, width=W)
    tb = s.shapes.add_textbox(PInches(0.8), PInches(5.4), W - PInches(1.6), PInches(1.6)); tf = tb.text_frame; tf.word_wrap = True
    r = tf.paragraphs[0].add_run(); r.text = "纸上花火 · Paper Fireworks"; pfont(r, SERIF, 40, SEAL, True)
    p = tf.add_paragraph(); r = p.add_run(); r.text = "一台走得进去的青春时光机 · 古早怀旧主题 VR 体验"; pfont(r, FONT, 18, IVORY)
    s.notes_slide.notes_text_frame.text = "大家好，我的策划是《纸上花火》。一句话——戴上头显，回到我们 QQ 空间、花火杂志的那个夏天。"

    slide("一句话概念", [
        "戴上头显，回到 2005–2012 那个 QQ 空间、非主流、花火杂志的年代，",
        "在一间被时光封存的少女卧室里，翻动每一件旧物，唤醒一段记忆，",
        "完成一次对青春的温柔告别。"],
        notes="它不是一个要赢的游戏，是一座可以走进去的、属于一代人的记忆博物馆。")
    slide("为什么做它 · 立意", [
        "母题：美好转瞬即逝；每个人的青春都值得被温柔收藏一次。",
        "情绪关键词：怀旧 · 治愈 · 沉浸 · 温柔的告别。",
        "市面怀旧多是图片与短视频，没人让你真正「走回去」——VR 能做到。"],
        notes="怀旧是这几年最强的情绪，但没人让你真正走回去。VR 能做到，这就是我的切入点。")
    slide("目标用户与市场", [
        "核心：90/95 后（QQ 空间原住民）；扩展：Y2K 美学党、治愈体验玩家。",
        "场景：个人头显 / 线下怀旧快闪展览 / 品牌联名（汽水·文具·音乐 App）。"],
        notes="这群人是消费主力又最吃怀旧；线下展览和联名是它天然的商业出口。")
    slide("世界观与场景", [
        "一间 12㎡ 千禧少女卧室：碎花床、海报墙、CRT 电脑+QQ、卡带机、",
        "梳妆台、铁皮文具盒、相片墙，窗外黄昏蝉鸣。",
        "用「道具密度」而非「场景大小」堆沉浸感。"], img=imgs["room"],
        notes="一进门，你就该脱口而出：这就是我那时候的房间。")
    slide("核心玩法", [
        "第一人称 VR 漫游 + 手柄交互。",
        "唤醒旧物 → 收集记忆碎片 → 集齐解锁结局。",
        "无失败、无倒计时、治愈向。"],
        notes="玩法很轻：自由地翻、看、听。每唤醒一件旧物，得到一片记忆。")
    slide("亮点 · 会说话的旧物（8 件 = 8 片记忆）", [
        "台式机 → QQ 空间（闪图、非主流签名、自动播放空间音乐）。",
        "卡带机 → 那年的歌；杂志 → 竖排言情；同学录 → 毕业留言。",
        "大头贴墙 → 合照；拍立得 → 拍照留念；日历 → 那年夏天的倒数。",
        "信纸 → 写一封寄不出的信。集齐 8 片记忆，解锁结局。"], img=imgs["qq"],
        notes="每件旧物都是一段记忆的钥匙。重点说两个：开机进 QQ 空间，闪图、非主流签名、自动播放的空间音乐全回来了；还有写一封寄不出的信。")
    slide("情感仪式 · 让怀旧变成参与", [
        "① 写一封寄不出的信 → 信化作光飘出窗外。",
        "② 拍立得留念 → 贴上相片墙。",
        "把「看」变成「参与」，情感更深。"],
        notes="光看会腻，所以我设计了两个主动仪式。写信，是把没说出口的话交给那个夏天；它不会寄出，但你放下了。")
    slide("美术风格", [
        "千禧 Y2K × 古早言情 × 非主流：暖黄夕照 + 玫红网格 + 衬线竖排。",
        "电影级怀旧后处理：暖调 + 暗角 + 胶片颗粒 + 泛黄。",
        "把「塑料 3D」压成「会动的旧照片」。"], img=imgs["mag"],
        notes="美术是命门。我用后处理把画面压成一张会动的旧照片，这是它和普通 3D 场景最大的区别。")
    slide("技术方案", [
        "Unity + URP（Bloom/调色/颗粒/暗角）。",
        "XR Interaction Toolkit（手柄射线 + 抓取）；TextMeshPro 竖排。",
        "UI 两模式：Overlay（系统/进度）+ WorldSpace（QQ/杂志贴物体）。",
        "防眩晕：瞬移 + 舒适转向 + 暗角；可出桌面演示版。"],
        notes="用主流 VR 管线，UI 两种模式都用上，专门做了防眩晕。")
    slide("叙事流程", [
        "黄昏入场 → 探索唤醒 → 集齐记忆（推到入夜）→ 写信 → 拍照 →",
        "房间褪色光化、旧物飘散 → 落款金句。"], img=imgs["end"],
        notes="结尾让整个房间在夕阳里慢慢褪色、化成光飘散——美好转瞬即逝，你完成了对青春的告别。")
    slide("音乐与音效", [
        "原创年代感主题曲（钢琴 + 卡带底噪）。",
        "环境音：蝉鸣 / 电扇 / 拨号上网 / QQ 提示音。",
        "全部原创或免版税，规避版权。"],
        notes="声音是怀旧的另一半。年代金曲用原创致敬版，规避版权又保留那个味道。")
    slide("商业与展览价值", [
        "怀旧 = 强情绪流量；每帧可截图传播。",
        "适配线下怀旧快闪展览 + 品牌联名；长尾话题强。"],
        notes="它不止是作业：怀旧自带传播力，线下展览和联名是清晰的落地路径。")
    # closing
    sc = prs.slides.add_slide(blank); sc.background.fill.solid(); sc.background.fill.fore_color.rgb = BG
    tb = sc.shapes.add_textbox(PInches(1), PInches(2.6), W - PInches(2), PInches(2)); tf = tb.text_frame; tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    r = tf.paragraphs[0].add_run(); r.text = "那年夏天，我们都以为来日方长。"; pfont(r, SERIF, 30, SEAL, True)
    p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = "纸上花火 · 谢谢观看"; pfont(r, FONT, 16, GOLD)
    sc.notes_slide.notes_text_frame.text = "《纸上花火》——把一代人的青春，温柔地收藏一次。谢谢大家。"

    path = os.path.join(OUT, "纸上花火_答辩PPT.pptx")
    prs.save(path); print("PPTX:", path, "| slides:", len(prs.slides._sldIdLst))

build_word()
build_ppt()
print("DONE")
