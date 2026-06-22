# -*- coding: utf-8 -*-
# 按 Gemini 视觉规范生成《纸上花火》答辩 PPT 终版
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

BASE=os.path.dirname(os.path.abspath(__file__)); OUT=os.path.join(BASE,"策划大作业_纸上花火"); IMG=os.path.join(OUT,"效果图")
IM={"room":os.path.join(IMG,"01_关键视觉_卧室黄昏.png"),"qq":os.path.join(IMG,"02_WorldSpaceUI_QQ空间.png"),
    "mag":os.path.join(IMG,"03_花火杂志_竖排金句.png"),"end":os.path.join(IMG,"04_结局_褪色光化.png")}

BG=RGBColor(0x16,0x10,0x1A); CARD=RGBColor(0xE8,0xDF,0xD3); WHITE=RGBColor(0xF5,0xEB,0xD8)
WHITE70=RGBColor(0xB9,0xB1,0xA4); GOLD=RGBColor(0xE8,0xC3,0x6A); ROSE=RGBColor(0xD9,0x6E,0x8C)
NIGHT=RGBColor(0x2C,0x3A,0x47); DARK=RGBColor(0x16,0x10,0x1A); DARK70=RGBColor(0x4a,0x42,0x4e)
SERIF="Noto Serif CJK SC"; SANS="Noto Sans CJK SC"; MONO="Courier New"

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
W,H=prs.slide_width,prs.slide_height; BLANK=prs.slide_layouts[6]

def newslide(notes=""):
    s=prs.slides.add_slide(BLANK)
    s.background.fill.solid(); s.background.fill.fore_color.rgb=BG
    if notes: s.notes_slide.notes_text_frame.text=notes
    return s
def setrun(r,font,size,color,bold=False,spc=None):
    f=r.font; f.size=Pt(size); f.bold=bold; f.color.rgb=color; f.name=font
    rPr=r._r.get_or_add_rPr()
    for tag in ('a:latin','a:ea','a:cs'):
        el=rPr.find(qn(tag))
        if el is None: el=rPr.makeelement(qn(tag),{}); rPr.append(el)
        el.set('typeface',font)
    if spc is not None: rPr.set('spc',str(int(spc*100)))
def alpha(shape,pct):
    sf=shape._element.spPr.find(qn('a:solidFill')); srgb=sf.find(qn('a:srgbClr'))
    srgb.append(srgb.makeelement(qn('a:alpha'),{'val':str(int(pct*1000))}))
def rect(s,l,t,w,h,color,a=None):
    sh=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,l,t,w,h); sh.fill.solid(); sh.fill.fore_color.rgb=color
    sh.line.fill.background(); sh.shadow.inherit=False
    if a is not None: alpha(sh,a)
    return sh
def fullimg(s,path):
    s.shapes.add_picture(path,0,0,width=W,height=H)
def cropimg(s,path,l,t,w,h,cl=0.0,cr=0.0):
    p=s.shapes.add_picture(path,l,t,width=w,height=h); p.crop_left=cl; p.crop_right=cr; return p
def block(s,l,t,w,h,lines,anchor=MSO_ANCHOR.TOP,align=PP_ALIGN.LEFT):
    tb=s.shapes.add_textbox(l,t,w,h); tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for i,ln in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=ln.get('align',align)
        p.space_after=Pt(ln.get('after',8)); p.space_before=Pt(ln.get('before',0))
        if ln.get('lh'): p.line_spacing=ln['lh']
        r=p.add_run(); r.text=ln['t']; setrun(r,ln.get('f',SANS),ln['s'],ln.get('c',WHITE70),ln.get('b',False),ln.get('spc'))
    return tb
def pagenum(s,i):
    block(s,W-Inches(2.2),Inches(0.3),Inches(1.9),Inches(0.4),
          [{'t':f"PAGE {i:02d} / 14",'f':MONO,'s':11,'c':ROSE,'spc':1,'align':PP_ALIGN.RIGHT}],align=PP_ALIGN.RIGHT)
def grid(s):
    for k in range(1,9):
        x=Emu(int(W*k/9)); c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,x,0,x,H)
        c.line.color.rgb=RGBColor(0xFF,0xFF,0xFF); c.line.width=Pt(0.5)
        ln=c.line._get_or_add_ln(); ln.find(qn('a:solidFill')).find(qn('a:srgbClr')).append(
            ln.makeelement(qn('a:alpha'),{'val':'8000'}))
    for k in range(1,6):
        y=Emu(int(H*k/6)); c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,0,y,W,y)
        c.line.color.rgb=RGBColor(0xFF,0xFF,0xFF); c.line.width=Pt(0.5)
        ln=c.line._get_or_add_ln(); ln.find(qn('a:solidFill')).find(qn('a:srgbClr')).append(
            ln.makeelement(qn('a:alpha'),{'val':'8000'}))
def corners(s,l,t,w,h):
    g=Inches(0.12)
    for ch,(x,y,ax) in {"⌜":(l-g,t-g,PP_ALIGN.LEFT),"⌝":(l+w-Inches(0.4)+g,t-g,PP_ALIGN.RIGHT),
                        "⌞":(l-g,t+h-Inches(0.5)+g,PP_ALIGN.LEFT),"⌟":(l+w-Inches(0.4)+g,t+h-Inches(0.5)+g,PP_ALIGN.RIGHT)}.items():
        block(s,x,y,Inches(0.4),Inches(0.5),[{'t':ch,'f':MONO,'s':26,'c':ROSE,'align':ax}])

# ---- LAYOUTS ----
def layA(title,bullets,notes,i,sub=None):
    s=newslide(notes); grid(s); pagenum(s,i)
    lines=[{'t':title,'f':SERIF,'s':36,'c':WHITE,'b':True,'after':14}]
    if sub: lines.append({'t':sub,'f':SANS,'s':20,'c':GOLD,'after':18})
    for b in bullets: lines.append({'t':"—  "+b,'f':SANS,'s':19,'c':WHITE70,'after':12,'lh':1.25})
    block(s,Inches(1.4),Inches(1.5),Inches(10.5),Inches(5),lines,anchor=MSO_ANCHOR.MIDDLE)
def layB(title,img,bullets,notes,i,cl=0.29,cr=0.29,corner=False):
    s=newslide(notes); pagenum(s,i)
    iw=Inches(5.6)
    cropimg(s,img,0,0,iw,H,cl,cr)
    if corner: corners(s,Inches(0.0),Inches(0.0),iw,H)
    lines=[{'t':title,'f':SERIF,'s':32,'c':WHITE,'b':True,'after':14}]
    for b in bullets: lines.append({'t':"—  "+b,'f':SANS,'s':18,'c':WHITE70,'after':11,'lh':1.25})
    block(s,iw+Inches(0.7),Inches(1.2),W-iw-Inches(1.3),Inches(5.1),lines,anchor=MSO_ANCHOR.MIDDLE)
def layC(title,bullets,notes,i,img=None):
    s=newslide(notes); pagenum(s,i)
    cl,ct,cw,chh=Inches(1.0),Inches(0.9),Inches(11.33),Inches(5.7)
    rect(s,cl,ct,cw,chh,CARD)
    if img:
        iw,ih=Inches(5.2),Inches(2.93)  # 16:9 不变形
        s.shapes.add_picture(img,cl+Inches(0.5),ct+(chh-ih)//2,width=iw,height=ih)
        tx=cl+Inches(6.0); tw=cw-Inches(6.6)
    else:
        tx=cl+Inches(0.9); tw=cw-Inches(1.8)
    lines=[{'t':title,'f':SERIF,'s':30,'c':DARK,'b':True,'after':14}]
    for b in bullets: lines.append({'t':"—  "+b,'f':SANS,'s':18,'c':DARK70,'after':11,'lh':1.25})
    block(s,tx,ct+Inches(0.5),tw,chh-Inches(1.0),lines,anchor=MSO_ANCHOR.MIDDLE)
def layFlow(title,flow,notes,i):
    s=newslide(notes); grid(s); pagenum(s,i)
    block(s,Inches(1.4),Inches(1.3),Inches(10),Inches(1),[{'t':title,'f':SERIF,'s':36,'c':WHITE,'b':True}])
    block(s,Inches(1.4),Inches(2.8),Inches(10.5),Inches(3),
          [{'t':flow,'f':SERIF,'s':26,'c':GOLD,'lh':2.0,'b':True}],anchor=MSO_ANCHOR.MIDDLE)

# ---- P1 cover ----
s=newslide("大家好，我的策划是《纸上花火》。一句话——戴上头显，回到我们 QQ 空间、花火杂志的那个夏天。")
fullimg(s,IM["room"]); rect(s,0,0,W,H,BG,a=22); rect(s,0,Inches(3.0),W,Inches(4.5),BG,a=62)
block(s,Inches(0.9),Inches(4.4),Inches(11),Inches(2.6),[
    {'t':"纸上花火",'f':SERIF,'s':54,'c':GOLD,'b':True,'after':4},
    {'t':"PAPER  FIREWORKS",'f':MONO,'s':18,'c':WHITE,'spc':2,'after':14},
    {'t':"一台走得进去的青春时光机 · 古早怀旧 VR 体验",'f':SANS,'s':20,'c':WHITE70}])

# ---- P2-P13 ----
layA("一句话概念",
    ["回到 2005–2012——QQ 空间、非主流、花火杂志的年代。",
     "在一间被时光封存的少女卧室里，翻动旧物、唤醒记忆，",
     "与青春做一次温柔的道别。"],
    "它不是一个要赢的游戏，是一座可以走进去的、属于一代人的记忆博物馆。",2)
layA("立意 · 为什么做它",
    ["母题：美好转瞬即逝；每个人的青春都值得被妥善收藏一次。",
     "定位：不是「通关」的游戏，而是一座可走入的时代记忆博物馆。",
     "关键词：怀旧 · 治愈 · 沉浸 · 告别。"],
    "怀旧是这几年最强的情绪，但没人让你真正走回去。VR 能做到，这就是切入点。",3)
layA("目标用户与市场",
    ["核心：90 / 95 后——QQ 空间原住民、花火读者、非主流亲历者。",
     "扩展：Y2K 美学党、情感治愈与沉浸体验玩家。",
     "场景：个人头显 / 线下怀旧快闪展览 / 品牌联名（汽水·文具·音乐 App）。"],
    "这群人是消费主力又最吃怀旧；线下展览和联名是它天然的商业出口。",4)
layB("世界观与场景",IM["room"],
    ["约 12㎡ 的千禧少女卧室：碎花床、海报墙、CRT 电脑、卡带机、相片墙。",
     "靠「道具密度」而非「场景大小」堆叠沉浸感。",
     "窗外天幕从黄昏过渡到入夜，驱动室内氛围。"],
    "一进门，你就该脱口而出：这就是我那时候的房间。",5,cl=0.18,cr=0.40)
layA("核心玩法",
    ["第一人称 VR 漫游 + 手柄交互。",
     "唤醒旧物 → 收集记忆碎片 → 集齐解锁结局。",
     "无失败、无倒计时、治愈向。"],
    "玩法很轻：自由地翻、看、听。每唤醒一件旧物，得到一片记忆。",6)
layB("会说话的旧物 · 8 件 = 8 片记忆",IM["qq"],
    ["台式机 → QQ 空间（闪图、非主流签名、自动播放空间音乐）。",
     "卡带机 → 那年的歌；杂志 → 竖排言情；同学录 → 毕业留言。",
     "信纸 → 写一封寄不出的信。集齐 8 片记忆，解锁结局。"],
    "每件旧物都是一段记忆的钥匙。开机进 QQ 空间，闪图、非主流签名、自动播放的空间音乐全回来了。",7,corner=True)
layC("情感仪式 · 让怀旧变成参与",
    ["① 写一封寄不出的信 → 信化作光飘出窗外。",
     "② 拍立得留念 → 贴上相片墙。",
     "把「看」变成「参与」，情感更深。"],
    "光看会腻，所以设计了两个主动仪式。写信，是把没说出口的话交给那个夏天。",8)
layC("美术风格",
    ["千禧 Y2K × 古早言情 × 非主流。",
     "暖黄夕照 + 玫红网格 + 衬线竖排。",
     "电影级怀旧后处理：把「塑料 3D」压成「会动的旧照片」。"],
    "美术是命门。我用后处理把画面压成一张会动的旧照片，这是它和普通 3D 场景最大的区别。",9,img=IM["mag"])
layC("技术方案",
    ["Unity + URP（Bloom / 调色 / 颗粒 / 暗角）。",
     "XR Interaction Toolkit（手柄射线 + 抓取）。",
     "UI 两模式：Overlay（系统/进度）+ WorldSpace（QQ/杂志贴物体）。",
     "防眩晕：瞬移 + 舒适转向 + 暗角；可出桌面演示版。"],
    "用主流 VR 管线，UI 两种模式都用上，专门做了防眩晕。",10)
layFlow("叙事流程",
    "黄昏入场  →  探索唤醒  →  集齐记忆  →  写信  →  拍照  →  褪色光化  →  落款金句",
    "结尾让整个房间在夕阳里慢慢褪色、化成光飘散——美好转瞬即逝，你完成了对青春的告别。",11)
layA("音乐与音效",
    ["原创年代主题曲（钢琴 + 卡带底噪）。",
     "环境音：蝉鸣 / 电扇 / 拨号上网 / QQ 提示音。",
     "全部原创或免版税，规避版权。"],
    "声音是怀旧的另一半。年代金曲用原创致敬版，规避版权又保留那个味道。",12)
layA("商业与展览价值",
    ["怀旧 = 强情绪流量；每帧画面可截图传播。",
     "适配线下怀旧快闪展览 + 品牌联名。",
     "长尾话题度与衍生传播潜力明显。"],
    "它不止是作业：怀旧自带传播力，线下展览和联名是清晰的落地路径。",13)

# ---- P14 ending ----
s=newslide("《纸上花火》——把一代人的青春，温柔地收藏一次。谢谢大家。")
fullimg(s,IM["end"]); rect(s,0,0,W,H,BG,a=45)
block(s,Inches(1.5),Inches(2.7),Inches(10.33),Inches(1.4),
    [{'t':"那年夏天，我们都以为来日方长。",'f':SERIF,'s':44,'c':GOLD,'b':True,'align':PP_ALIGN.CENTER}],align=PP_ALIGN.CENTER)
c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(5.2),Inches(4.35),Inches(8.13),Inches(4.35))
c.line.color.rgb=ROSE; c.line.width=Pt(1)
block(s,Inches(1.5),Inches(4.6),Inches(10.33),Inches(0.8),
    [{'t':"纸上花火 · 谢谢观看",'f':SANS,'s':20,'c':WHITE,'spc':2,'align':PP_ALIGN.CENTER}],align=PP_ALIGN.CENTER)

path=os.path.join(OUT,"纸上花火_答辩PPT_终版(Gemini视觉).pptx")
prs.save(path); print("SAVED:",path,"| slides:",len(prs.slides._sldIdLst))
