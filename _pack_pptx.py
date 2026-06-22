import os
from pptx import Presentation
from pptx.util import Inches
OUT="策划大作业_纸上花火"
notes=[
"大家好，我的策划是《纸上花火》——戴上头显，回到我们 QQ 空间、花火杂志的那个夏天。",
"它不是一个要赢的游戏，是一座可以走进去的、属于一代人的记忆博物馆。",
"怀旧是这几年最强的情绪，但没人让你真正走回去。VR 能做到，这是切入点。",
"这群人是消费主力又最吃怀旧；线下展览和联名是天然的商业出口。",
"一进门就该脱口而出：这就是我那时候的房间。靠道具密度而非场景大小堆沉浸。",
"玩法很轻：自由地翻、看、听。每唤醒一件旧物，得到一片记忆。",
"每件旧物都是一段记忆的钥匙。开机进 QQ 空间，闪图、签名、空间音乐全回来了；重点还有写一封寄不出的信。",
"光看会腻，所以设计了两个主动仪式。写信，是把没说出口的话交给那个夏天。",
"美术是命门。后处理把画面压成一张会动的旧照片，这是和普通 3D 场景最大的区别。",
"用主流 VR 管线，UI 两种模式都用上，专门做了防眩晕。",
"结尾让房间在夕阳里褪色、化成光飘散——美好转瞬即逝，完成对青春的告别。",
"声音是怀旧的另一半。年代金曲用原创致敬版，规避版权又保留那个味道。",
"它不止是作业：怀旧自带传播力，线下展览和联名是清晰的落地路径。",
"《纸上花火》——把一代人的青春，温柔地收藏一次。谢谢大家。",
]
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
blank=prs.slide_layouts[6]
for i in range(1,15):
    s=prs.slides.add_slide(blank)
    s.shapes.add_picture(f"/tmp/deck/s{i:02d}.png",0,0,width=prs.slide_width,height=prs.slide_height)
    s.notes_slide.notes_text_frame.text=notes[i-1]
p=os.path.join(OUT,"纸上花火_答辩PPT_终版(Gemini视觉).pptx")
prs.save(p); print("SAVED",p,"slides",len(prs.slides._sldIdLst))
