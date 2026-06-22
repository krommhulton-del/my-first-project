# -*- coding: utf-8 -*-
# 用 DS 润色稿重排最终 Word（甲方需求 + 乙方解决方案，嵌入 4 张效果图）
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = os.path.dirname(os.path.abspath(__file__))
OUT  = os.path.join(BASE, "策划大作业_纸上花火")
IMG  = os.path.join(OUT, "效果图")
imgs = [os.path.join(IMG, n) for n in
        ["01_关键视觉_卧室黄昏.png","02_WorldSpaceUI_QQ空间.png","03_花火杂志_竖排金句.png","04_结局_褪色光化.png"]]
caps = ["▲ 效果图 1 · 关键视觉 / 卧室黄昏","▲ 效果图 2 · WorldSpace UI / 开机进入 QQ 空间",
        "▲ 效果图 3 · 互动旧物 / 花火杂志翻页与竖排","▲ 效果图 4 · 结局 / 房间褪色光化与告别"]
FONT="微软雅黑"; SERIF="思源宋体"
WINE=RGBColor(0x7A,0x2A,0x35); NIGHT=RGBColor(0x2C,0x3A,0x47); GREY=RGBColor(0x88,0x78,0x68)

def cjk(run,name=FONT,size=10.5,bold=False,color=None):
    run.font.name=name; run.font.size=Pt(size); run.font.bold=bold
    if color: run.font.color.rgb=color
    rPr=run._element.get_or_add_rPr(); rf=rPr.find(qn('w:rFonts'))
    if rf is None: rf=rPr.makeelement(qn('w:rFonts'),{}); rPr.insert(0,rf)
    rf.set(qn('w:eastAsia'),name)
def para(doc,t,size=10.5,bold=False,color=None,align=None,after=4,name=FONT):
    p=doc.add_paragraph()
    if align: p.alignment=align
    p.paragraph_format.space_after=Pt(after)
    cjk(p.add_run(t),name,size,bold,color); return p
def H(doc,t,size=14.5): para(doc,t,size=size,bold=True,color=WINE,after=6,name=SERIF)
def bl(doc,items,size=10.5):
    for it in items:
        p=doc.add_paragraph(); p.paragraph_format.left_indent=Pt(14); p.paragraph_format.space_after=Pt(2)
        cjk(p.add_run("· "+it),FONT,size)
def tbl(doc,rows):
    t=doc.add_table(rows=len(rows),cols=len(rows[0])); t.style='Table Grid'
    for i,row in enumerate(rows):
        for j,c in enumerate(row):
            cell=t.cell(i,j); cell.paragraphs[0].clear(); cjk(cell.paragraphs[0].add_run(c),FONT,9.5,bold=(i==0))
    doc.add_paragraph()

doc=Document()
para(doc,"《纸上花火 · Paper Fireworks》",24,True,WINE,WD_ALIGN_PARAGRAPH.CENTER,2,SERIF)
para(doc,"古早怀旧主题 · 沉浸式 VR 体验 ｜ 策划大作业（含 甲方需求文档 + 乙方解决方案）",11,False,GREY,WD_ALIGN_PARAGRAPH.CENTER,2)
para(doc,"（自己一组）",10,False,RGBColor(0xA0,0x90,0x80),WD_ALIGN_PARAGRAPH.CENTER,12)

# ===== 第一部分 需求文档（甲方）=====
para(doc,"第一部分　需求文档（甲方）",16,True,NIGHT,after=8,name=SERIF)
H(doc,"一、项目概述")
tbl(doc,[["项","内容"],
 ["产品名称","《纸上花火 · Paper Fireworks》"],
 ["产品类型","古早怀旧主题 · 沉浸式 VR 体验 / 叙事探索轻游戏"],
 ["目标平台","PC VR（Meta Quest / Pico 串流），另提供桌面体验版"],
 ["体验时长","单次 15–25 分钟，可重复进入"],
 ["一句话概念","戴上头显，回到 2005–2012——QQ 空间、非主流、花火杂志的年代。一间被时光封存的少女卧室里，翻动旧物，唤醒一段集体记忆，与青春做一次温柔的道别。"]])
H(doc,"二、立意与情感内核")
bl(doc,["母题：美好转瞬即逝；每个人的青春，都值得被妥善收藏一次。",
 "定位：它并非以「通关」为目标的游戏，而是一座可走入的时代记忆博物馆。",
 "体验关键词：怀旧 · 治愈 · 沉浸 · 告别。"])
H(doc,"三、目标用户与使用场景")
bl(doc,["核心人群：90 后 / 95 后——QQ 空间原住民、花火与飞言情读者、非主流时代的亲历者。",
 "扩展人群：Y2K 千禧美学爱好者、情感治愈类与沉浸体验类玩家、文创及展览观众。",
 "使用场景：个人头显体验 / 线下怀旧主题快闪展览 / 品牌联名合作（汽水、文具、音乐 App）。"])
H(doc,"四、需求功能模块（我方提出，不设可行性与成本限制）")
para(doc,"4.1　场景",bold=True,size=11)
para(doc,"高度还原 2005–2012 年间一间典型的千禧少女卧室。具体要素包括：碎花床品、海报与明星贴纸墙面、CRT 台式电脑（运行 QQ）、卡带机 / MP3、梳妆台、铁皮文具盒、毛绒玩具、星空夜光顶贴、铁丝相片墙、成摞花火杂志。窗外为夏日傍晚景象——晚霞、蝉鸣与转动的电风扇。")
para(doc,"4.2　可交互旧物（共 8 件，每件对应一片记忆，集齐即解锁结局）",bold=True,size=11)
tbl(doc,[["旧物","交互方式","唤起的记忆"],
 ["老式台式机","开机后点击进入 QQ 空间（闪图、非主流签名、自动播放的空间音乐、踩一踩）","那年的网络青春"],
 ["卡带机 / MP3","放入磁带或戴上耳机，播放年代歌曲（原创致敬版本，规避版权）","一首歌，一段心事"],
 ["花火 / 飞言情杂志","抓取并翻页，呈现古早言情插画与竖排金句","课桌下偷看的言情"],
 ["同学录","翻开，展示手写体留言「给十年后的你」","毕业季"],
 ["大头贴墙","拼贴大头贴的小互动","和好朋友的合照"],
 ["钢笔与信纸","手柄书写或语音，写一封「寄不出的信」，可封存或点燃","没说出口的话"],
 ["拍立得相机","取景按下快门，吐出做旧照片并贴上相片墙","把此刻收藏起来"],
 ["日历 / 闹钟","撕下日历或按停闹钟，推进「那年夏天」的时间线","倒数着的暑假"]])
para(doc,"4.3　核心玩法",bold=True,size=11)
para(doc,"第一人称 VR 漫游 + 手柄交互；唤醒旧物 → 收集记忆碎片 → 集齐后解锁结局；写信与拍照为两个情感仪式环节；无失败条件、无倒计时，整体基调为治愈向。")
para(doc,"4.4　UI 形态（两种并存）",bold=True,size=11)
bl(doc,["Overlay 模式：承载系统信息、收集进度、相册；",
 "World Space 模式：QQ 空间界面、杂志内页等直接附着于场景物体表面。"])
para(doc,"4.5　音乐与音效：卡带底噪、蝉鸣、电扇转动声、键盘敲击声、QQ 提示音、拨号上网音，搭配一首原创年代主题曲。",size=10.5)
para(doc,"4.6　粒子与氛围：夕照中的浮尘、飘动的窗帘、入夜后的萤火或远处烟花、结局阶段的褪色光化效果。",size=10.5)
para(doc,"4.7　结局：集齐全部记忆后，房间在夕阳中逐步褪色、光化、飘散，最终浮现落款金句。",size=10.5)
H(doc,"五、核心体验流程")
para(doc,"进入房间（黄昏）→ 自由探索，逐件唤醒旧物 → 集齐记忆碎片 → 写一封寄不出的信 → 拍一张留念照 → 房间褪色光化，完成告别 → 浮现落款金句。",size=10.5)
H(doc,"六、验收与效果期望")
bl(doc,["还原度：进入场景后第一反应是「这就是我当年的房间」。",
 "情感共鸣：至少有一件旧物能让人产生明确的情绪触动。",
 "沉浸感与舒适度：无明显眩晕感，用户愿意长时间停留在场景中。",
 "传播性：每帧画面具备截图传播的质感，适合展览与联名场合使用。"])

doc.add_page_break()
# ===== 第二部分 解决方案（乙方）=====
para(doc,"第二部分　解决方案（乙方）",16,True,NIGHT,after=8,name=SERIF)
H(doc,"一、对需求的理解")
para(doc,"甲方需要的并非传统意义上的「游戏」，而是一台情绪时光机。一切设计的核心目标在于：让用户进入场景后，立刻辨认出自己的青春，并产生留恋。还原度、声音质感、光影与文字的处理，是决定体验成败的关键。")
H(doc,"二、设计总览")
tbl(doc,[["甲方需求","对应解决方案"],
 ["高度还原千禧卧室","3D 建模 + PBR 做旧材质 + 电影级怀旧后处理（暖调、暗角、颗粒感、泛黄倾向），将「塑料感 3D」压制成「旧照片」质感"],
 ["旧物承载记忆","每件旧物作为可交互体（XR 抓取 / 触发），触发后播放一段小型演出（音乐、动画、竖排文字），并生成一片记忆碎片"],
 ["自由漫游且不晕眩","基于 XR Interaction Toolkit；采用瞬移移动、舒适转向与暗角减晕机制；无失败条件，无倒计时"],
 ["两种 UI 形态","Overlay 承载系统信息、收集进度与相册；World Space 将 QQ 空间界面、杂志内页等直接置于场景物体之上"],
 ["怀旧音乐且规避版权","采用原创致敬风格 BGM，搭配真实环境音素材，全面规避版权风险"],
 ["治愈向告别体验","结局粒子演出：房间逐层褪色 → 光化 → 飘散，最终浮现落款金句"]])
H(doc,"三、美术与场景方案")
bl(doc,["风格定位：千禧 Y2K × 古早言情 × 非主流。视觉关键词：暖黄夕照、玫红网格、闪光渐变、衬线字体竖排。",
 "场景规模：约 12㎡ 的单场景空间，依靠道具密度而非面积来堆叠沉浸感。",
 "光照变化：窗外天幕从黄昏过渡至入夜，驱动室内氛围演变。",
 "色彩体系：暖黄 #E8C36A / 玫红 #D96E8C / 米白 #E8DFD3 / 夜蓝 #2C3A47 / 高光白 #F5EBD8，整体统一压旧处理。"])
H(doc,"四、互动设计方案（XR 手柄交互）")
bl(doc,["台式机：按下电源键 → 射线点击 QQ 图标 → 弹出 World Space 形态的 QQ 空间界面。",
 "卡带机：抓取磁带放入卡槽 / 戴上耳机 → 播放年代歌曲，同步展示竖排歌词。",
 "杂志：抓取并翻页 → 展示插画与竖排古早言情金句。",
 "信纸与钢笔：手柄模拟书写或语音输入 → 可选择封口或点燃 → 信纸化作光点或灰烬飘出窗外。",
 "拍立得：举起取景 → 按下快门 → 吐出做旧照片 → 可粘贴至相片墙。"])
H(doc,"五、系统与玩法方案")
bl(doc,["记忆收集系统：场景内 8 件旧物对应 8 片记忆，全部收集后触发结局。",
 "情感仪式系统：写信与拍照两个环节，结果存入相册，可供回看。",
 "软性引导：未唤醒的旧物带有微弱高光与轻柔声响，引导玩家靠近。",
 "无失败设计：不设任何失败条件与时间压力。"])
H(doc,"六、叙事流程（节拍设计）")
para(doc,"黄昏时分进入房间 → 软引导推动探索，逐件唤醒旧物 → 集齐 8 片记忆（推进至入夜）→ 写信 → 拍照 → 房间褪色、光化，旧物逐一飘散 → 黑场浮现落款：「那年夏天，我们都以为来日方长。」",size=10.5)
H(doc,"七、UI 方案（双模式并行）")
bl(doc,["Overlay 模式：承载入场标题卡、记忆收集进度（8 颗星）、相册回看、设置项。",
 "World Space 模式：CRT 屏幕内的 QQ 空间页面、杂志与同学录页面、旧物悬浮提示文字。"])
H(doc,"八、音频方案")
bl(doc,["原创主题曲：钢琴主导，叠加卡带底噪质感。","环境层：蝉鸣、电扇转动、拨号音、QQ 提示音等。",
 "交互音效：与每件旧物的触发行为绑定。","版权策略：全部采用原创或免版税素材。"])
H(doc,"九、技术方案")
bl(doc,["引擎：Unity + URP（启用 Bloom、暖调校色、颗粒感、暗角等后处理效果）。",
 "交互框架：XR Interaction Toolkit（Ray + Direct 交互），搭配 XR Device Simulator 用于桌面调试。",
 "文字渲染：TextMeshPro 支持竖排排版与中文衬线字体。",
 "输出形态：PC VR 包 + 桌面演示版 + 360° 全景截图。"])
H(doc,"十、防眩晕与舒适体验措施")
bl(doc,["移动方式：瞬移位移 + 舒适转向 + 移动期间触发暗角效果。",
 "交互区域：可交互物体布局在视线舒适区与手边范围内。",
 "画面控制：无快速强制位移，无高频率闪烁。","使用姿态：支持坐姿与站姿两种体验模式。"])
H(doc,"十一、效果呈现参考图（概念效果图）")
for path,cap in zip(imgs,caps):
    if os.path.exists(path):
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(path,width=Inches(5.6))
        para(doc,cap,9,False,GREY,WD_ALIGN_PARAGRAPH.CENTER,10)
H(doc,"十二、开发里程碑")
tbl(doc,[["阶段","内容"],
 ["P0 原型阶段","单房间场景 + 基础漫游 + 1 件可交互旧物"],
 ["P1 美术阶段","全道具建模 + 做旧处理 + 整体调色"],
 ["P2 交互阶段","全部 8 件旧物交互 + 记忆收集系统 + 写信与拍照功能"],
 ["P3 音频与叙事阶段","主题曲与背景音乐 + 环境音 + 完整叙事流程 + 结局演出"],
 ["P4 打磨阶段","防眩晕优化 + 桌面版适配 + 录屏与展览物料准备"]])
H(doc,"十三、风险与应对")
tbl(doc,[["风险","应对"],
 ["年代音乐版权","采用原创致敬版本与免版税素材，全程规避商用版权问题"],
 ["VR 眩晕","瞬移、舒适转向、暗角三重机制叠加控制"],
 ["还原度不足","收集大量真实年代道具参考素材，引入目标年龄段用户参与测试与反馈"],
 ["体验偏静态","通过写信与拍照两个主动性仪式环节，提升参与感与情感投入度"]])
H(doc,"十四、商业与展览价值")
para(doc,"怀旧本身即具有强烈情绪流量。本作天然适配线下快闪展览与品牌联名场景（汽水、文具、音乐 App 等），画面每帧具备截图传播属性，长尾话题度与衍生传播潜力明显。",size=10.5)

path=os.path.join(OUT,"纸上花火_策划大作业_终稿(DS润色).docx")
doc.save(path); print("SAVED:",path)
