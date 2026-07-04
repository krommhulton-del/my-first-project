#!/usr/bin/env python3
"""简历生成器 v0.1 — 结构化 YAML → 单页精美 PDF + 可编辑 docx

用法:
    python scripts/gen_resume.py input/简历.yaml -o output/

产出(同一份 YAML,两种格式):
    - <name>_简历.pdf   排版精良的投递版(playwright 打印,矢量清晰)
    - <name>_简历.docx  可编辑版(python-docx,方便客户自己微调)

铁律 2:只排版你在 YAML 里填的真实内容,绝不替任何人编造经历。
    脚本不生成、不润色事实,只做版式。

YAML 字段(见 input/简历示例.yaml):
    name, title, contact{phone,email,city,extra}, summary,
    experience[{company,role,period,bullets[]}],
    education[{school,major,period,extra}], skills[], projects[{name,desc}]
"""
import argparse
import os
import sys
from html import escape
from pathlib import Path

import yaml

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from playwright.sync_api import sync_playwright

INK = "#16243D"
ACCENT = "#16243D"
MUTED = "#5A6472"
RULE = "#D8D3C8"

REPO_ROOT = Path(__file__).resolve().parent.parent


# ---------- 共用:找 chromium ----------
def _chromium_executable():
    env_bin = os.environ.get("CHROME_EXECUTABLE") or os.environ.get("PLAYWRIGHT_CHROMIUM_EXECUTABLE")
    if env_bin and Path(env_bin).exists():
        return env_bin
    for base in filter(None, [os.environ.get("PLAYWRIGHT_BROWSERS_PATH"), "/opt/pw-browsers"]):
        base_p = Path(base)
        for c in [base_p / "chromium", *sorted(base_p.glob("chromium-*/chrome-linux/chrome"), reverse=True)]:
            if c.exists():
                return str(c)
    return None


# ---------- HTML → PDF ----------
def build_html(cv: dict) -> str:
    name = escape(str(cv.get("name", "")))
    title = escape(str(cv.get("title", "")))
    c = cv.get("contact", {}) or {}
    contact_bits = [c.get("phone"), c.get("email"), c.get("city")]
    contact_line = " · ".join(escape(str(x)) for x in contact_bits if x)
    extra = escape(str(c.get("extra", ""))) if c.get("extra") else ""

    def section(titletext, inner):
        return (f'<section><h2>{escape(titletext)}</h2>{inner}</section>') if inner else ""

    # 个人简介
    summary_html = ""
    if cv.get("summary"):
        summary_html = section("个人简介", f'<p class="summary">{escape(str(cv["summary"]).strip())}</p>')

    # 工作经历
    exp_items = ""
    for e in cv.get("experience", []) or []:
        bullets = "".join(f"<li>{escape(str(b))}</li>" for b in (e.get("bullets") or []))
        exp_items += (
            '<div class="entry">'
            f'<div class="entry-head"><span class="org">{escape(str(e.get("company","")))}'
            f'<span class="role"> · {escape(str(e.get("role","")))}</span></span>'
            f'<span class="period">{escape(str(e.get("period","")))}</span></div>'
            f'<ul>{bullets}</ul></div>'
        )
    exp_html = section("工作经历", exp_items)

    # 教育背景
    edu_items = ""
    for e in cv.get("education", []) or []:
        extra_e = f' <span class="edu-extra">{escape(str(e.get("extra")))}</span>' if e.get("extra") else ""
        edu_items += (
            '<div class="entry compact">'
            f'<div class="entry-head"><span class="org">{escape(str(e.get("school","")))}'
            f'<span class="role"> · {escape(str(e.get("major","")))}</span></span>'
            f'<span class="period">{escape(str(e.get("period","")))}</span></div>'
            f'{("<div class=sub>"+extra_e+"</div>") if extra_e else ""}</div>'
        )
    edu_html = section("教育背景", edu_items)

    # 技能
    skills_html = ""
    if cv.get("skills"):
        lis = "".join(f"<li>{escape(str(s))}</li>" for s in cv["skills"])
        skills_html = section("技能标签", f'<ul class="skills">{lis}</ul>')

    # 项目
    proj_items = ""
    for p in cv.get("projects", []) or []:
        proj_items += (
            '<div class="entry compact">'
            f'<div class="entry-head"><span class="org">{escape(str(p.get("name","")))}</span></div>'
            f'<div class="sub">{escape(str(p.get("desc","")))}</div></div>'
        )
    proj_html = section("项目经历", proj_items)

    return f"""<!DOCTYPE html><html lang="zh-CN"><head><meta charset="UTF-8"><style>
  @page {{ size: A4; margin: 0; }}
  * {{ margin:0; padding:0; box-sizing:border-box; -webkit-print-color-adjust:exact; print-color-adjust:exact; }}
  body {{ width:210mm; min-height:297mm; padding:16mm 16mm 14mm;
         font-family:"Noto Sans CJK SC","Noto Sans SC","PingFang SC","Microsoft YaHei",sans-serif;
         color:#1c2530; font-size:10.5pt; line-height:1.5; }}
  header {{ border-bottom:2px solid {ACCENT}; padding-bottom:10px; margin-bottom:6px; }}
  .name {{ font-size:24pt; font-weight:800; color:{INK}; letter-spacing:1px; }}
  .title {{ font-size:12pt; color:{ACCENT}; font-weight:600; margin-top:2px; }}
  .contact {{ font-size:9pt; color:{MUTED}; margin-top:6px; }}
  .contact .extra {{ display:block; margin-top:2px; }}
  section {{ margin-top:13px; }}
  h2 {{ font-size:11pt; color:{INK}; font-weight:800; letter-spacing:2px;
       padding-left:9px; border-left:4px solid {ACCENT}; margin-bottom:7px; }}
  .summary {{ color:#33404f; }}
  .entry {{ margin-bottom:9px; }}
  .entry.compact {{ margin-bottom:6px; }}
  .entry-head {{ display:flex; justify-content:space-between; align-items:baseline; }}
  .org {{ font-weight:700; color:{INK}; font-size:10.8pt; }}
  .role {{ font-weight:600; color:{ACCENT}; }}
  .period {{ font-size:9pt; color:{MUTED}; white-space:nowrap; padding-left:10px; }}
  ul {{ list-style:none; margin-top:3px; }}
  li {{ position:relative; padding-left:14px; margin:2px 0; color:#33404f; }}
  li::before {{ content:""; position:absolute; left:2px; top:8px; width:4px; height:4px;
              background:{ACCENT}; border-radius:50%; }}
  .sub {{ color:{MUTED}; font-size:9.5pt; margin-top:1px; }}
  .edu-extra {{ color:{MUTED}; }}
  ul.skills {{ display:flex; flex-wrap:wrap; gap:6px 10px; }}
  ul.skills li {{ padding:3px 12px; background:#F2F0EB; border-radius:4px; color:{INK}; }}
  ul.skills li::before {{ display:none; }}
</style></head><body>
  <header>
    <div class="name">{name}</div>
    <div class="title">{title}</div>
    <div class="contact">{contact_line}{f'<span class="extra">{extra}</span>' if extra else ''}</div>
  </header>
  {summary_html}{exp_html}{edu_html}{skills_html}{proj_html}
</body></html>"""


def render_pdf(cv: dict, out_path: Path):
    html = build_html(cv)
    exe = _chromium_executable()
    kwargs = {"executable_path": exe} if exe else {}
    with sync_playwright() as p:
        browser = p.chromium.launch(**kwargs)
        try:
            page = browser.new_page()
            page.set_content(html, wait_until="networkidle")
            page.pdf(path=str(out_path), format="A4", print_background=True,
                     margin={"top": "0", "bottom": "0", "left": "0", "right": "0"})
        finally:
            browser.close()


# ---------- docx ----------
def _set_cn_font(run, size, color=(0x1C, 0x25, 0x30), bold=False, cn="微软雅黑", en="Segoe UI"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)
    run.font.name = en
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), cn)


def _heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    _set_cn_font(p.add_run(text), 12, color=(0x16, 0x24, 0x3D), bold=True)
    return p


def render_docx(cv: dict, out_path: Path):
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Cm(1.6)
        section.left_margin = section.right_margin = Cm(1.8)

    # 头部
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    _set_cn_font(p.add_run(str(cv.get("name", ""))), 22, color=(0x16, 0x24, 0x3D), bold=True)
    if cv.get("title"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        _set_cn_font(p.add_run(str(cv["title"])), 12, color=(0x16, 0x24, 0x3D), bold=True)
    c = cv.get("contact", {}) or {}
    contact_line = " · ".join(str(x) for x in [c.get("phone"), c.get("email"), c.get("city")] if x)
    if contact_line:
        p = doc.add_paragraph()
        _set_cn_font(p.add_run(contact_line), 9.5, color=(0x5A, 0x64, 0x72))
    if c.get("extra"):
        p = doc.add_paragraph()
        _set_cn_font(p.add_run(str(c["extra"])), 9.5, color=(0x5A, 0x64, 0x72))

    if cv.get("summary"):
        _heading(doc, "个人简介")
        p = doc.add_paragraph()
        _set_cn_font(p.add_run(str(cv["summary"]).strip()), 10.5)

    if cv.get("experience"):
        _heading(doc, "工作经历")
        for e in cv["experience"]:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            _set_cn_font(p.add_run(str(e.get("company", ""))), 11, color=(0x16, 0x24, 0x3D), bold=True)
            if e.get("role"):
                _set_cn_font(p.add_run(f" · {e['role']}"), 10.5, color=(0x16, 0x24, 0x3D), bold=True)
            if e.get("period"):
                _set_cn_font(p.add_run(f"    {e['period']}"), 9, color=(0x5A, 0x64, 0x72))
            for b in e.get("bullets", []) or []:
                bp = doc.add_paragraph(style="List Bullet")
                bp.paragraph_format.space_after = Pt(1)
                _set_cn_font(bp.add_run(str(b)), 10.5, color=(0x33, 0x40, 0x4F))

    if cv.get("education"):
        _heading(doc, "教育背景")
        for e in cv["education"]:
            p = doc.add_paragraph()
            _set_cn_font(p.add_run(str(e.get("school", ""))), 11, color=(0x16, 0x24, 0x3D), bold=True)
            if e.get("major"):
                _set_cn_font(p.add_run(f" · {e['major']}"), 10.5, color=(0x16, 0x24, 0x3D), bold=True)
            if e.get("period"):
                _set_cn_font(p.add_run(f"    {e['period']}"), 9, color=(0x5A, 0x64, 0x72))
            if e.get("extra"):
                sp = doc.add_paragraph()
                _set_cn_font(sp.add_run(str(e["extra"])), 9.5, color=(0x5A, 0x64, 0x72))

    if cv.get("skills"):
        _heading(doc, "技能标签")
        p = doc.add_paragraph()
        _set_cn_font(p.add_run("  |  ".join(str(s) for s in cv["skills"])), 10.5)

    if cv.get("projects"):
        _heading(doc, "项目经历")
        for pr in cv["projects"]:
            p = doc.add_paragraph()
            _set_cn_font(p.add_run(str(pr.get("name", ""))), 11, color=(0x16, 0x24, 0x3D), bold=True)
            if pr.get("desc"):
                sp = doc.add_paragraph()
                _set_cn_font(sp.add_run(str(pr["desc"])), 10.5, color=(0x33, 0x40, 0x4F))

    doc.save(out_path)


def main():
    ap = argparse.ArgumentParser(description="简历生成器:YAML → 单页 PDF + docx")
    ap.add_argument("input", type=Path, help="简历 YAML 文件")
    ap.add_argument("-o", "--out", type=Path, default=Path("output"), help="输出目录")
    ap.add_argument("--no-pdf", action="store_true", help="只出 docx")
    args = ap.parse_args()

    if not args.input.exists():
        sys.exit(f"错误: 找不到 YAML {args.input}")
    cv = yaml.safe_load(args.input.read_text(encoding="utf-8"))
    if not isinstance(cv, dict) or not cv.get("name"):
        sys.exit("错误: YAML 需为对象且至少包含 name 字段")

    args.out.mkdir(parents=True, exist_ok=True)
    stem = str(cv["name"])
    docx_path = args.out / f"{stem}_简历.docx"
    render_docx(cv, docx_path)
    print(f"完成 docx → {docx_path}")

    if not args.no_pdf:
        pdf_path = args.out / f"{stem}_简历.pdf"
        render_pdf(cv, pdf_path)
        print(f"完成 PDF  → {pdf_path}")

    print("提示:PDF 用于投递,docx 可自己微调;脚本只排版,不改事实。")


if __name__ == "__main__":
    main()
